#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Convert all markdown files in docs/ to docx
- Copies docs/ to markdown/ as backup (per user request)
- Converts each .md in docs/ to .docx alongside
- Also creates docx/ folder with same structure for easy download
"""
import os
import re
import shutil
from pathlib import Path

ROOT = Path("/home/user/liveforever")
DOCS = ROOT / "docs"
BACKUP_MD = ROOT / "markdown"  # new markdown folder at root
DOCX_OUT = ROOT / "docx"  # consolidated docx folder

# Ensure backup
if BACKUP_MD.exists():
    shutil.rmtree(BACKUP_MD)
shutil.copytree(DOCS, BACKUP_MD)
print(f"Copied {DOCS} -> {BACKUP_MD}")

# Also copy to docs/markdown as alternative location per ambiguous instruction
DOCS_MARKDOWN = DOCS / "markdown"
if DOCS_MARKDOWN.exists():
    shutil.rmtree(DOCS_MARKDOWN)
# Copy only md files to docs/markdown to avoid recursion
DOCS_MARKDOWN.mkdir(parents=True, exist_ok=True)
for md_file in DOCS.rglob("*.md"):
    # skip if inside docs/markdown itself (avoid copy loop)
    if DOCS_MARKDOWN in md_file.parents:
        continue
    rel = md_file.relative_to(DOCS)
    dest = DOCS_MARKDOWN / rel
    dest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(md_file, dest)
print(f"Copied md files to {DOCS_MARKDOWN}")

# Clean docx output
if DOCX_OUT.exists():
    shutil.rmtree(DOCX_OUT)
DOCX_OUT.mkdir(parents=True, exist_ok=True)

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

def parse_inline(text, paragraph):
    """Parse **bold** and normal text into runs"""
    # Split by bold markers
    # Pattern **text**
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) >= 4:
            inner = part[2:-2]
            run = paragraph.add_run(inner)
            run.bold = True
        else:
            # Handle also *italic*? Keep simple
            # Also handle `code`
            # Remove markdown link syntax [text](url) -> text (url)
            # Process links
            # We'll replace [xxx](yyy) -> xxx (yyy)
            link_match = re.sub(r'\[([^\]]+)\]\(([^\)]+)\)', r'\1 (\2)', part)
            paragraph.add_run(link_match)
    return paragraph

def md_to_docx(md_path: Path, docx_path: Path):
    doc = Document()
    # Set default font for Chinese support - use a common font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Microsoft YaHei'
    font.size = Pt(11)
    # Try to set east asia font
    try:
        r = style._element
        # python-docx doesn't easily set eastAsia, but ok
    except:
        pass

    with open(md_path, 'r', encoding='utf-8', errors='ignore') as f:
        content = f.read()

    lines = content.splitlines()
    i = 0
    in_table = False
    table_rows = []

    def flush_table():
        nonlocal table_rows, in_table
        if not table_rows:
            return
        # Parse table rows: split by |
        # Filter out separator rows like |---|---|
        filtered = []
        for row in table_rows:
            # Check if separator row
            if re.match(r'^\s*\|?\s*[-|:\s]+\s*\|?\s*$', row):
                # if contains only - | : and spaces, skip if it's separator
                # But ensure it doesn't contain alphanumeric Chinese
                if re.search(r'[a-zA-Z0-9\u4e00-\u9fff]', row.replace('-','').replace('|','').replace(':','').replace(' ','')) is None:
                    continue
            filtered.append(row)
        if not filtered:
            table_rows = []
            in_table = False
            return
        # Build docx table
        # Determine max columns
        parsed = []
        for row in filtered:
            # Trim leading/trailing |
            row = row.strip()
            if row.startswith('|'):
                row = row[1:]
            if row.endswith('|'):
                row = row[:-1]
            cells = [c.strip() for c in row.split('|')]
            parsed.append(cells)
        if not parsed:
            table_rows = []
            in_table = False
            return
        max_cols = max(len(r) for r in parsed)
        # Normalize
        for r in parsed:
            while len(r) < max_cols:
                r.append('')
        # Create table
        table = doc.add_table(rows=len(parsed), cols=max_cols)
        table.style = 'Light Grid Accent 1'
        for r_idx, row in enumerate(parsed):
            for c_idx, cell_text in enumerate(row):
                cell = table.cell(r_idx, c_idx)
                cell.text = re.sub(r'\*\*(.*?)\*\*', r'\1', cell_text)  # strip bold for table
                # Make header bold
                if r_idx == 0:
                    for p in cell.paragraphs:
                        for run in p.runs:
                            run.bold = True
        table_rows = []
        in_table = False
        doc.add_paragraph('')  # spacing

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Detect table row
        if '|' in line and (line.count('|') >= 1):
            # Heuristic: if next line is separator or this line looks like table, start table
            # Look ahead: if this line contains | and next line is separator, it's table start
            # Or if we're already in table and line has |
            next_is_sep = False
            if i+1 < len(lines):
                nxt = lines[i+1].strip()
                if re.match(r'^\|?[\s\-\:\|]+\|?$', nxt) and '-' in nxt:
                    next_is_sep = True
            if next_is_sep or in_table or (stripped.startswith('|') and stripped.endswith('|')):
                if not in_table:
                    in_table = True
                    table_rows = []
                table_rows.append(line)
                i += 1
                continue
            # else not table, fall through

        if in_table:
            # If current line doesn't look like table, flush
            if '|' not in line or stripped == '':
                flush_table()
            else:
                # still table but maybe blank?
                table_rows.append(line)
                i += 1
                continue

        if stripped == '':
            # Empty line -> add empty paragraph for spacing, but not too many
            doc.add_paragraph('')
            i += 1
            continue

        # Headings
        m = re.match(r'^(#{1,6})\s+(.*)', line)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            # Remove markdown bold inside heading for title
            clean = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
            clean = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', clean)
            if level == 1:
                p = doc.add_heading(clean, level=1)
            elif level == 2:
                p = doc.add_heading(clean, level=2)
            elif level == 3:
                p = doc.add_heading(clean, level=3)
            else:
                p = doc.add_heading(clean, level=4)
            i += 1
            continue

        # Horizontal rule
        if re.match(r'^[-*_]{3,}\s*$', stripped):
            p = doc.add_paragraph('─' * 40)
            p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            i += 1
            continue

        # Blockquote >
        if stripped.startswith('>'):
            bq_text = stripped.lstrip('> ').strip()
            # Collect consecutive blockquote lines
            bq_lines = [bq_text]
            j = i+1
            while j < len(lines) and lines[j].strip().startswith('>'):
                bq_lines.append(lines[j].strip().lstrip('> ').strip())
                j += 1
            full = '\n'.join(bq_lines)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(20)
            # Add with italic
            run = p.add_run(full)
            run.italic = True
            run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            # parse bold inside?
            # For simplicity, re-parse
            p.clear()
            parse_inline(full, p)
            for run in p.runs:
                run.italic = True
            i = j
            continue

        # List items
        list_match = re.match(r'^(\s*)([-*+]|\d+\.)\s+(.*)', line)
        if list_match:
            indent = len(list_match.group(1))
            marker = list_match.group(2)
            content_text = list_match.group(3)
            # Create paragraph with bullet/number
            p = doc.add_paragraph(style='List Bullet' if marker in ['-','*','+'] else 'List Number')
            p.paragraph_format.left_indent = Pt(20 + indent*10)
            parse_inline(content_text, p)
            i += 1
            # Handle continuation lines that are indented
            # (simple)
            continue

        # Normal paragraph - collect until blank or special
        para_lines = [line]
        j = i+1
        while j < len(lines):
            nl = lines[j]
            if nl.strip() == '':
                break
            if re.match(r'^(#{1,6})\s+', nl):
                break
            if re.match(r'^(\s*)([-*+]|\d+\.)\s+', nl):
                break
            if nl.strip().startswith('>'):
                break
            if '|' in nl and nl.count('|') >= 2:
                # potential table start
                if j+1 < len(lines) and re.match(r'^\|?[\s\-\:\|]+\|?$', lines[j+1].strip()):
                    break
            para_lines.append(nl)
            j += 1
        full_para = ' '.join([l.strip() for l in para_lines])
        p = doc.add_paragraph()
        parse_inline(full_para, p)
        i = j

    # Final flush if table left
    if in_table:
        flush_table()

    # Ensure parent dir exists
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))
    print(f"Converted {md_path} -> {docx_path}")

# Walk docs for md files
md_files = list(DOCS.rglob("*.md"))
# Exclude the markdown backup folder inside docs to avoid infinite loop
md_files = [p for p in md_files if DOCS_MARKDOWN not in p.parents]

print(f"Found {len(md_files)} markdown files in docs/")

for md_file in md_files:
    # Skip if file is inside docs/markdown (we already filtered)
    rel = md_file.relative_to(DOCS)
    # docx alongside original
    docx_alongside = md_file.with_suffix('.docx')
    # docx in consolidated folder
    docx_consolidated = DOCX_OUT / rel.with_suffix('.docx')
    
    try:
        md_to_docx(md_file, docx_alongside)
        # Also copy to consolidated, but if alongside is same as consolidated parent, copy
        if docx_alongside != docx_consolidated:
            docx_consolidated.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(docx_alongside, docx_consolidated)
    except Exception as e:
        print(f"FAILED {md_file}: {e}")
        import traceback
        traceback.print_exc()

print("Done")
# List outputs
print("\n=== Backup markdown folder ===")
for p in BACKUP_MD.rglob("*"):
    print(p.relative_to(ROOT))
print("\n=== Docx alongside in docs/ ===")
for p in DOCS.rglob("*.docx"):
    print(p.relative_to(ROOT))
print("\n=== Consolidated docx/ ===")
for p in DOCX_OUT.rglob("*.docx"):
    print(p.relative_to(ROOT))
